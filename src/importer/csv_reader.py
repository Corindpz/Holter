import csv
from pathlib import Path
from typing import List, Dict, Any
import chardet
from src.models import Ticket


def read_csv(path: str, mapping: Dict[str, str]) -> List[Dict[str, Any]]:
    file_path = Path(path)
    if file_path.suffix.lower() == ".xlsx":
        return _read_xlsx(file_path, mapping)
    return _read_csv(file_path, mapping)


def _read_csv(file_path: Path, mapping: Dict[str, str]) -> List[Dict[str, Any]]:
    raw_bytes = file_path.read_bytes()
    encoding = chardet.detect(raw_bytes)["encoding"] or "utf-8"
    text = raw_bytes.decode(encoding, errors="replace")
    reader = csv.DictReader(text.splitlines())
    rows = []
    for row in reader:
        normalized = {}
        for internal_key, sf_column in mapping.items():
            normalized[internal_key] = row.get(sf_column, "").strip()
        rows.append(normalized)
    return rows


def _read_xlsx(file_path: Path, mapping: Dict[str, str]) -> List[Dict[str, Any]]:
    import openpyxl
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    ws = wb.active
    headers = [str(cell.value).strip() if cell.value is not None else "" for cell in next(ws.iter_rows(min_row=1, max_row=1))]
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        raw_row = {headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)}
        normalized = {}
        for internal_key, sf_column in mapping.items():
            normalized[internal_key] = raw_row.get(sf_column, "")
        rows.append(normalized)
    wb.close()
    return rows


def normalize_ticket(raw: Dict[str, Any], semaine_code: str) -> Ticket:
    return Ticket(
        id=raw.get("id", ""),
        semaine_code=semaine_code,
        objet=raw.get("objet") or None,
        priorite=raw.get("priorite") or None,
        statut=raw.get("statut") or None,
        produit=raw.get("produit") or None,
        site=raw.get("site") or None,
        description=raw.get("description") or None,
        date_creation=raw.get("date_creation") or None,
    )
