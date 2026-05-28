import io
import json
from datetime import datetime, timezone
from fastapi import APIRouter, HTTPException
from fastapi.responses import Response

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from src.db.connection import get_db
from src.export.pdf_generator import generate_pdf
from src.review.review_manager import ReviewManager
from src.config import get_settings

router = APIRouter(tags=["export"])
_rm = ReviewManager()

# ── Styles Excel ─────────────────────────────────────────────
_BG_HEADER  = "1a1a1a"
_BG_MV      = "3d1212"
_BG_IV      = "3d2e00"
_BG_SECU    = "0d1f3d"
_BG_CLOS    = "0d2e1e"
_BG_WATCH   = "1e1030"
_FG_WHITE   = "EDEDE9"
_FG_MV      = "E24B4A"
_FG_IV      = "F59E0B"
_FG_SECU    = "3B82F6"
_FG_CLOS    = "10B981"
_FG_WATCH   = "8B5CF6"

_THIN = Side(style="thin", color="333333")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _hdr(ws, row: int, col: int, value: str, width: int = 18) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(bold=True, color=_FG_WHITE, name="Calibri", size=10)
    cell.fill = PatternFill("solid", fgColor=_BG_HEADER)
    cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
    cell.border = _BORDER
    ws.column_dimensions[get_column_letter(col)].width = width


def _cell(ws, row: int, col: int, value, fg: str = _FG_WHITE, bold: bool = False) -> None:
    cell = ws.cell(row=row, column=col, value=value)
    cell.font = Font(color=fg, name="Calibri", size=10, bold=bold)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    cell.border = _BORDER


def _signal_fg(signal: str | None) -> str:
    return {"MV": _FG_MV, "IV": _FG_IV, "SECU": _FG_SECU}.get(signal or "", _FG_WHITE)


def _decision_fg(decision: str | None) -> str:
    return {
        "ANALYSE_REQUISE": _FG_MV,
        "SURVEILLER": _FG_WATCH,
        "CLOS": _FG_CLOS,
    }.get(decision or "", _FG_WHITE)


def _row_fill(decision: str | None) -> str | None:
    return {
        "ANALYSE_REQUISE": _BG_MV,
        "SURVEILLER": _BG_WATCH,
        "CLOS": None,
    }.get(decision or "")


def generate_excel(rows: list, semaine_code: str, expert_nom: str | None) -> bytes:
    wb = openpyxl.Workbook()

    # ── Feuille 1 : Résumé ────────────────────────────────────
    ws_sum = wb.active
    ws_sum.title = "Résumé"
    ws_sum.sheet_view.showGridLines = False

    kpis = {"total": 0, "analyse_requise": 0, "surveiller": 0, "clos": 0, "sans_analyse": 0}
    for r in rows:
        kpis["total"] += 1
        dec = (r.get("decision") or "").upper()
        if dec == "ANALYSE_REQUISE":   kpis["analyse_requise"] += 1
        elif dec == "SURVEILLER":      kpis["surveiller"] += 1
        elif dec == "CLOS":            kpis["clos"] += 1
        else:                          kpis["sans_analyse"] += 1

    summary_data = [
        ("Semaine",            semaine_code),
        ("Expert",             expert_nom or "—"),
        ("Généré le",          datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("",                   ""),
        ("Total tickets",      kpis["total"]),
        ("Analyse requise",    kpis["analyse_requise"]),
        ("Surveiller",         kpis["surveiller"]),
        ("Clos",               kpis["clos"]),
        ("Non analysés",       kpis["sans_analyse"]),
    ]

    for i, (label, value) in enumerate(summary_data, start=2):
        c_label = ws_sum.cell(row=i, column=2, value=label)
        c_label.font = Font(bold=True, color=_FG_WHITE, name="Calibri", size=11)
        c_val = ws_sum.cell(row=i, column=3, value=value)
        c_val.font = Font(color=_FG_WHITE, name="Calibri", size=11)

    ws_sum.column_dimensions["B"].width = 22
    ws_sum.column_dimensions["C"].width = 28
    ws_sum.sheet_properties.tabColor = "E24B4A"

    # ── Feuille 2 : Tous les tickets ──────────────────────────
    ws = wb.create_sheet("Tickets")
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = "A2"

    headers = [
        ("ID",          14), ("Date",         14), ("Objet",        38),
        ("Produit",     22), ("Site",          22), ("Statut SF",    14),
        ("Priorité",    12), ("Décision IA",   18), ("Signal",       10),
        ("Niveau",      12), ("Confiance",     12), ("Pass",          8),
        ("Action expert",14),("Décision finale",18),("Commentaire",  30),
        ("Mots-clés",   28), ("Articles cités",28),
    ]
    for col, (label, width) in enumerate(headers, start=1):
        _hdr(ws, 1, col, label, width)
    ws.row_dimensions[1].height = 22

    for r_idx, r in enumerate(rows, start=2):
        dec = r.get("decision") or ""
        bg = _row_fill(dec)
        fill = PatternFill("solid", fgColor=bg) if bg else None
        row_data = [
            r.get("id", ""),
            r.get("date_creation", "")[:10] if r.get("date_creation") else "",
            r.get("objet", "") or "",
            r.get("produit", "") or "",
            r.get("site", "") or "",
            r.get("statut", "") or "",
            r.get("priorite", "") or "",
            dec,
            r.get("signal", "") or "",
            r.get("niveau", "") or "",
            f"{round(r['confiance'] * 100)}%" if r.get("confiance") is not None else "",
            str(r.get("passe_finale", "")) if r.get("passe_finale") is not None else "",
            r.get("action_expert", "") or "",
            r.get("decision_finale", "") or "",
            r.get("commentaire", "") or "",
            ", ".join(json.loads(r["mots_cles"]) if r.get("mots_cles") else []),
            ", ".join(json.loads(r["articles_cites"]) if r.get("articles_cites") else []),
        ]
        for col, val in enumerate(row_data, start=1):
            cell = ws.cell(row=r_idx, column=col, value=val)
            cell.font = Font(color=_FG_WHITE, name="Calibri", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (3, 15, 16, 17)))
            cell.border = _BORDER
            if fill:
                cell.fill = fill
            # Coloration par décision/signal sur les colonnes clés
            if col == 8:
                cell.font = Font(color=_decision_fg(dec), name="Calibri", size=10, bold=True)
            elif col == 9:
                cell.font = Font(color=_signal_fg(r.get("signal")), name="Calibri", size=10, bold=True)
        ws.row_dimensions[r_idx].height = 18

    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"
    ws.sheet_properties.tabColor = "3B82F6"

    # ── Feuille 3 : Signaux vigilance ─────────────────────────
    ws_sig = wb.create_sheet("Signaux vigilance")
    ws_sig.sheet_view.showGridLines = False
    ws_sig.freeze_panes = "A2"

    sig_headers = [
        ("ID", 14), ("Date", 14), ("Objet", 38), ("Produit", 22),
        ("Site", 22), ("Signal", 10), ("Niveau", 12), ("Décision", 18),
        ("Confiance", 12), ("Raisonnement", 55), ("Articles cités", 30),
    ]
    for col, (label, width) in enumerate(sig_headers, start=1):
        _hdr(ws_sig, 1, col, label, width)
    ws_sig.row_dimensions[1].height = 22

    sig_rows = [r for r in rows if r.get("signal") and r.get("decision") != "CLOS"]
    for r_idx, r in enumerate(sig_rows, start=2):
        sig = r.get("signal") or ""
        sig_data = [
            r.get("id", ""),
            r.get("date_creation", "")[:10] if r.get("date_creation") else "",
            r.get("objet", "") or "",
            r.get("produit", "") or "",
            r.get("site", "") or "",
            sig,
            r.get("niveau", "") or "",
            r.get("decision", "") or "",
            f"{round(r['confiance'] * 100)}%" if r.get("confiance") is not None else "",
            r.get("raisonnement", "") or "",
            ", ".join(json.loads(r["articles_cites"]) if r.get("articles_cites") else []),
        ]
        for col, val in enumerate(sig_data, start=1):
            cell = ws_sig.cell(row=r_idx, column=col, value=val)
            cell.font = Font(color=_FG_WHITE, name="Calibri", size=10)
            cell.alignment = Alignment(vertical="top", wrap_text=(col in (3, 10, 11)))
            cell.border = _BORDER
            if col == 6:
                cell.font = Font(color=_signal_fg(sig), name="Calibri", size=10, bold=True)
            elif col == 8:
                cell.font = Font(color=_decision_fg(r.get("decision")), name="Calibri", size=10, bold=True)
        ws_sig.row_dimensions[r_idx].height = 32

    ws_sig.auto_filter.ref = f"A1:{get_column_letter(len(sig_headers))}1"
    ws_sig.sheet_properties.tabColor = "E24B4A"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.post("/export/pdf/{semaine_code}")
async def export_pdf(semaine_code: str):
    if _rm.pending_count(semaine_code) > 0:
        raise HTTPException(
            status_code=409,
            detail="Export verrouillé : des tickets ANALYSE_REQUISE n'ont pas encore été validés."
        )
    with get_db() as conn:
        sem = conn.execute("SELECT * FROM semaines WHERE code = ?", (semaine_code,)).fetchone()
        if not sem:
            raise HTTPException(status_code=404, detail="Semaine non trouvée")
        rows = conn.execute(
            """SELECT t.id, t.objet, t.priorite, t.produit, t.site, t.statut,
                      a.decision, a.signal, a.niveau, a.confiance,
                      a.raisonnement, a.articles_cites, a.capa_suggere,
                      d.action_expert, d.commentaire, d.horodatage
               FROM tickets t
               LEFT JOIN analyses a ON t.id = a.ticket_id AND t.semaine_code = a.semaine_code
               LEFT JOIN decisions d ON t.id = d.ticket_id AND t.semaine_code = d.semaine_code
               WHERE t.semaine_code = ?""",
            (semaine_code,),
        ).fetchall()
        dict_count = conn.execute("SELECT COUNT(*) as n FROM dictionnaire").fetchone()["n"]

    tickets = []
    kpis = {"total": 0, "analyse_requise": 0, "surveiller": 0, "clos": 0}
    for r in rows:
        t = dict(r)
        raw_articles = t.get("articles_cites")
        t["articles_cites"] = json.loads(raw_articles) if raw_articles else []
        tickets.append(t)
        kpis["total"] += 1
        dec = (t.get("decision") or "CLOS").upper()
        if dec == "ANALYSE_REQUISE":
            kpis["analyse_requise"] += 1
        elif dec == "SURVEILLER":
            kpis["surveiller"] += 1
        else:
            kpis["clos"] += 1

    settings = get_settings()
    data = {
        "semaine_code": semaine_code,
        "date_debut": sem["date_debut"],
        "date_fin": sem["date_fin"],
        "expert_nom": sem["expert_nom"],
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "model_used": settings.get("model_override") or "auto-detect",
        "kpis": kpis,
        "tickets": tickets,
        "dictionary_version": dict_count,
        "threshold_clos": settings["confidence_threshold_clos"],
    }
    pdf_bytes = generate_pdf(data)
    filename = f"holter_pms_{semaine_code}.pdf"
    return Response(
        content=pdf_bytes,
        media_type="application/pdf",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )


@router.get("/export/excel/{semaine_code}")
async def export_excel(semaine_code: str):
    with get_db() as conn:
        sem = conn.execute("SELECT * FROM semaines WHERE code = ?", (semaine_code,)).fetchone()
        if not sem:
            raise HTTPException(status_code=404, detail="Semaine non trouvée")
        rows = conn.execute(
            """SELECT t.id, t.objet, t.priorite, t.produit, t.site, t.statut,
                      t.date_creation,
                      a.decision, a.signal, a.niveau, a.confiance,
                      a.raisonnement, a.articles_cites, a.mots_cles, a.passe_finale,
                      d.action_expert, d.decision_finale, d.commentaire
               FROM tickets t
               LEFT JOIN analyses a ON t.id = a.ticket_id AND t.semaine_code = a.semaine_code
               LEFT JOIN decisions d ON t.id = d.ticket_id AND t.semaine_code = d.semaine_code
               WHERE t.semaine_code = ?
               ORDER BY a.confiance ASC NULLS LAST""",
            (semaine_code,),
        ).fetchall()

    excel_bytes = generate_excel(
        [dict(r) for r in rows],
        semaine_code,
        sem["expert_nom"],
    )
    filename = f"holter_pms_{semaine_code}.xlsx"
    return Response(
        content=excel_bytes,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename={filename}"},
    )
